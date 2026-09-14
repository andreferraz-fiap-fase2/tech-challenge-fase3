"""Recorte municipal dos indicadores educacionais históricos do Inep."""

import math
import re
from io import BytesIO
from typing import cast

import pandas as pd
from openpyxl import load_workbook

from src.domain.contract import ContractError
from src.domain.education_study import EDUCATION_FEATURES

INDICATORS = {
    "ATU": (EDUCATION_FEATURES[0], 0.0, 200.0),
    "DSU": (EDUCATION_FEATURES[1], 0.0, 100.0),
    "HAD": (EDUCATION_FEATURES[2], 0.0, 24.0),
}
KEYS = ["id_municipio", "rede_nome"]
_UF_CODES = dict(
    zip(
        "11 12 13 14 15 16 17 21 22 23 24 25 26 27 28 29 31 32 33 35 41 42 43 50 51 52 53".split(),
        "RO AC AM RR PA AP TO MA PI CE RN PB PE AL SE BA MG ES RJ SP PR SC RS MS MT GO DF".split(),
        strict=True,
    )
)
_REQUIRED = {
    "NU_ANO_CENSO",
    "SG_UF",
    "CO_MUNICIPIO",
    "NO_CATEGORIA",
    "NO_DEPENDENCIA",
    "FUN_AI_CAT_0",
}
_DEPENDENCIES = {"Total", "Federal", "Estadual", "Municipal", "Privada", "Pública"}


def parse_education_xlsx(payload: bytes, indicator: str) -> pd.DataFrame:
    """Lê bytes XLSX de 2021; seleciona Anos Iniciais, localização Total e duas redes."""
    if indicator not in INDICATORS:
        raise ContractError(f"Indicador educacional desconhecido: {indicator}")
    feature, lower, upper = INDICATORS[indicator]
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["MUNICIPIO"]:
            raise ContractError("Planilha Inep deve conter exclusivamente a aba MUNICIPIO")
        rows = workbook["MUNICIPIO"].iter_rows(values_only=True)
        columns: list[str] | None = None
        for position, row in enumerate(rows):
            if row and row[0] == "NU_ANO_CENSO":
                columns = [str(value).strip() for value in row]
                break
            if position >= 19:
                break
        if columns is None or not _REQUIRED.issubset(columns):
            raise ContractError("Cabeçalho Inep sem ano, município, rede ou Anos Iniciais")
        if any(columns.count(name) != 1 for name in _REQUIRED):
            raise ContractError("Cabeçalho Inep com colunas obrigatórias duplicadas")
        positions = {name: columns.index(name) for name in _REQUIRED}
        records: list[dict[str, str | float]] = []
        for row in rows:
            values = cast(tuple[object, ...], row)
            if not values or values[0] is None:
                continue
            if isinstance(values[0], str) and values[0].startswith(("Fonte:", "Nota:", "Notas:")):
                continue
            year = values[positions["NU_ANO_CENSO"]]
            if str(year).strip() != "2021":
                raise ContractError(f"Ano do indicador Inep={year}; esperado 2021")
            raw_code = values[positions["CO_MUNICIPIO"]]
            code = str(raw_code).strip()
            if isinstance(raw_code, float) and math.isfinite(raw_code) and raw_code.is_integer():
                code = str(int(raw_code))
            if not re.fullmatch(r"\d{7}", code) or code[:2] not in _UF_CODES:
                raise ContractError(f"Código municipal Inep inválido: {raw_code}")
            if str(values[positions["SG_UF"]]).strip() != _UF_CODES[code[:2]]:
                raise ContractError(f"UF incompatível com o código municipal Inep {code}")
            location = str(values[positions["NO_CATEGORIA"]]).strip()
            network = str(values[positions["NO_DEPENDENCIA"]]).strip()
            if location not in {"Total", "Urbana", "Rural"} or network not in _DEPENDENCIES:
                raise ContractError(f"Localização/rede Inep desconhecida: {location}/{network}")
            if location != "Total" or network not in {"Estadual", "Municipal"}:
                continue
            value = _parse_value(values[positions["FUN_AI_CAT_0"]], lower, upper, indicator)
            records.append({"id_municipio": code, "rede_nome": network, feature: value})
        if not records:
            raise ContractError("Recorte municipal Inep vazio para as redes Estadual/Municipal")
        frame = pd.DataFrame.from_records(records).astype({"id_municipio": "string"})
        if frame.duplicated(KEYS).any():
            raise ContractError(f"Chave município/rede duplicada no indicador {indicator}")
        return frame.sort_values(KEYS).reset_index(drop=True)
    finally:
        workbook.close()


def _parse_value(value: object, lower: float, upper: float, indicator: str) -> float:
    if value is None or (isinstance(value, str) and value.strip() in {"", "--"}):
        return float("nan")
    try:
        numeric = float(str(value).strip().replace(",", "."))
    except (TypeError, ValueError) as exc:
        raise ContractError(f"Valor não numérico no indicador {indicator}: {value}") from exc
    if not math.isfinite(numeric) or not lower <= numeric <= upper:
        raise ContractError(f"Valor fora do intervalo [{lower}, {upper}] em {indicator}: {value}")
    if indicator in {"ATU", "HAD"} and numeric == 0:
        raise ContractError(f"Indicador {indicator} deve ser positivo quando disponível")
    return numeric


def combine_education_context(frames: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Une as três fontes por município/rede e preserva ausências como NaN."""
    if set(frames) != set(INDICATORS):
        raise ContractError("Contexto educacional exige exatamente ATU, DSU e HAD")
    combined: pd.DataFrame | None = None
    for indicator in INDICATORS:
        frame = frames[indicator]
        feature = INDICATORS[indicator][0]
        if set(frame.columns) != set(KEYS + [feature]):
            raise ContractError(f"Colunas inesperadas no recorte do indicador {indicator}")
        if frame[KEYS].isna().any().any() or frame.duplicated(KEYS).any():
            raise ContractError(f"Chave município/rede ausente ou duplicada em {indicator}")
        combined = (
            frame.copy()
            if combined is None
            else combined.merge(frame, on=KEYS, how="outer", validate="one_to_one")
        )
    assert combined is not None
    return combined.sort_values(KEYS).reset_index(drop=True)
