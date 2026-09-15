"""Autor: André Mohallem Ferraz. Slides de decisão e roteiro externo de preparação."""

import json
from pathlib import Path

from documentos_word import AUTHOR, GOLD, GRAY, NAVY, ROOT, WordDocuments
from pptx import Presentation
from pptx.dml.color import RGBColor as PColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide
from pptx.util import Inches as PInches
from pptx.util import Pt as PPt

REPOSITORY = "https://github.com/andreferraz-fiap-fase2/tech-challenge-fase3"

NARRATIONS = [
    "Este projeto de André Mohallem Ferraz estima a probabilidade de alfabetização de um aluno do segundo ano, considerando seu contexto. O resultado observado segue o critério de setecentos e quarenta e três pontos de proficiência. A estimativa apoia perguntas sobre fatores associados, territórios, semelhanças regionais e metas, sem substituir a avaliação pedagógica individual.",
    "A base reúne avaliações reais em dois ciclos. Reconstruímos a Gold por aluno a partir da Silver e dos originais da Fase dois, com enriquecimento do IBGE. A nota define o alvo e fica fora dos preditores. Usamos três divisões por município em dois mil e vinte e três, com pré-processamento aprendido no treino, e teste separado em dois mil e vinte e quatro.",
    "A exploração encontrou quarenta e um vírgula sessenta e um por cento de não alfabetização e contextos repetidos, justificando validação municipal. No Inep, as medianas são dezenove vírgula cinco alunos por turma, noventa e quatro vírgula quatro por cento de funções docentes com superior e quatro vírgula três horas diárias. Cada histograma conta uma vez cada município e rede, sem multiplicar contextos pelos alunos.",
    "Comparamos baseline, regressão logística e Gradient Boosting. O modelo escolhido alcançou average precision de zero vírgula cinco um seis, contra zero vírgula quatro zero dois da referência. Essa métrica mede ordenação do risco. O limiar F dois recupera noventa e nove vírgula trinta por cento dos casos, mas sinaliza noventa e seis vírgula oitenta e quatro por cento dos alunos. A seletividade é baixa para triagem autônoma.",
    "O estudo complementar acrescenta três indicadores do Inep aos seis atributos originais, mantendo os mesmos alunos, divisões e parâmetros. A cobertura supera noventa e nove vírgula noventa e oito por cento. O ganho de average precision foi pequeno, positivo em duas divisões e negativo em uma. Sem novo teste independente, essa expansão não substitui o modelo final.",
    "Sobre os fatores associados, funções docentes com superior apresentam associação positiva com alfabetização; alunos por turma, negativa; horas diárias, próxima de zero. Isso não demonstra impacto causal. A influência no modelo responde outra pergunta: ao permutar atributos na validação, a maior perda de desempenho ocorre com a unidade federativa. Serviços públicos, população e PIB têm contribuições menores. Correlação contextual e importância preditiva não são a mesma medida.",
    "Entre municípios com pelo menos cem avaliações, Aracaju e Nossa Senhora do Socorro apresentam os maiores riscos médios previstos: aproximadamente sessenta e oito e sessenta e sete por cento. Os dez primeiros estão em Sergipe, evidenciando dependência estadual. Nos perfis econômicos de dois mil e vinte e três, Centro-Oeste e Sul são os mais próximos. Isso não implica taxas de alfabetização iguais nem constitui agrupamento automático de alunos.",
    "Para metas, mil quinhentos e noventa e dois municípios ficam abaixo da referência de dois mil e vinte e quatro. No cenário de oitenta por cento, são dois mil setecentos e sessenta e sete. Comparamos médias ponderadas das probabilidades mantendo a composição observada. Esses cenários não preveem dois mil e trinta nem estimam a probabilidade de descumprimento. Uma previsão futura exige novos ciclos e avaliação independente.",
    "A demonstração para o perfil histórico de Belo Horizonte municipal estima cinquenta e oito vírgula sessenta e seis por cento de alfabetização. A referência de cinquenta por cento classifica como alfabetizado; a política F dois sinaliza atenção. A probabilidade é a mesma, com decisões diferentes. O exemplo não é uma previsão individual para dois mil e vinte e seis.",
    "A recomendação é investigar territórios combinando risco, volume, cobertura e evidências pedagógicas locais. Para avançar em metas futuras, precisamos ampliar ciclos, obter atributos anteriores ao período previsto e reservar avaliação independente. O repositório indicado reúne código, métodos e resultados verificáveis. A contribuição é apoiar decisões com evidências e limites claros, distinguindo associação, previsão e causalidade.",
]

SCRIPT_WINDOWS = [
    "00:00–00:25",
    "00:25–00:55",
    "00:55–01:25",
    "01:25–01:55",
    "01:55–02:25",
    "02:25–03:00",
    "03:00–03:35",
    "03:35–04:05",
    "04:05–04:30",
    "04:30–05:00",
]


def video_script(version: str) -> str:
    """Roteiro externo; o PowerPoint contém notas vazias e nenhum áudio incorporado."""
    script = (
        "# Roteiro de preparação do vídeo executivo — Fase 3\n\n**Autor: "
        + AUTHOR
        + f" · Revisão {version} · Duração planejada: até 5 minutos**\n\n"
        + "As janelas somam 5 minutos e incluem pausas e transições. São uma estimativa "
        + "para ensaio, não uma duração de gravação medida. Material de apoio, fora do ZIP "
        + "e da pasta oficial. O vídeo oficial deverá ser gravado com a voz do autor. "
        + "As notas do PowerPoint permanecem vazias.\n\n"
    )
    for index, narration in enumerate(NARRATIONS, 1):
        script += f"## {SCRIPT_WINDOWS[index - 1]} — Slide {index}\n\n{narration}\n\n"
    return script + (
        "## Orientação de apresentação\n\n"
        "Ensaiar como reunião executiva e concluir em até 5 minutos. As cinco perguntas "
        "estratégicas são respondidas nos slides 6 a 8: fatores associados, influência no "
        "modelo, municípios de maior risco, semelhanças regionais e metas futuras. "
        "Distinguir associação de causalidade, semelhança contextual de agrupamento e "
        "cenário de previsão futura. Explicar AP sem confundir com acurácia. "
        "Não afirmar que o cenário de 80% prevê 2030.\n\n"
        f"[Repositório do projeto]({REPOSITORY}).\n"
    )


class SlideDocuments:
    def __init__(self, output: Path, version: str = "1.0") -> None:
        self.output = output
        self.version = WordDocuments(output, version).version
        output.mkdir(parents=True, exist_ok=True)

    def textbox(
        self,
        slide: Slide,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        size: int = 24,
        color: str = NAVY,
        bold: bool = False,
        link: str | None = None,
    ) -> None:
        shape = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
        shape.text_frame.word_wrap = True
        for index, line in enumerate(text.split("\n")):
            paragraph = (
                shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
            )
            paragraph.text = line
            paragraph.font.name = "Aptos"
            paragraph.font.size = PPt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = PColor.from_string(color)
            paragraph.space_after = PPt(8)
            if link:
                for run in paragraph.runs:
                    run.hyperlink.address = link
                    run.font.underline = True

    def wide(
        self, slide: Slide, text: str, y: float = 5.25, size: int = 20, color: str = GRAY
    ) -> None:
        self.textbox(slide, text, 0.65, y, 12, 1.35, size, color)

    def slide_base(self, presentation: PresentationType, title: str, number: int) -> Slide:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PInches(13.333), PInches(0.16))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PColor.from_string(GOLD)
        bar.line.fill.background()
        self.textbox(
            slide,
            "FIAP  /  INTELIGÊNCIA ANALÍTICA PARA ALFABETIZAÇÃO",
            0.55,
            0.3,
            12,
            0.3,
            11,
            GOLD,
            True,
        )
        self.textbox(slide, title, 0.55, 0.9, 12.2, 1.15, 31, NAVY, True)
        self.textbox(
            slide, AUTHOR + "  ·  Fase 3  ·  Setembro de 2026", 0.55, 7.05, 7.5, 0.3, 10, GRAY
        )
        self.textbox(
            slide,
            "GitHub · tech-challenge-fase3",
            8.4,
            7.03,
            3.4,
            0.3,
            10,
            GOLD,
            link=REPOSITORY,
        )
        self.textbox(slide, f"{number:02d}", 12.1, 7.03, 0.6, 0.3, 11, GOLD, True)
        slide.notes_slide.notes_text_frame.text = ""
        return slide

    def cards(self, slide: Slide, items: list[tuple[str, str]], y: float = 2.35) -> None:
        width = 11.8 / len(items)
        for index, (headline, body) in enumerate(items):
            x = 0.6 + index * width
            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                PInches(x),
                PInches(y),
                PInches(width - 0.15),
                PInches(2.6),
            )
            box.fill.solid()
            box.fill.fore_color.rgb = PColor.from_string("F1F4F6")
            box.line.fill.background()
            self.textbox(slide, headline, x + 0.22, y + 0.3, width - 0.6, 0.85, 32, NAVY, True)
            self.textbox(slide, body, x + 0.22, y + 1.2, width - 0.6, 1.15, 20, GRAY)

    def presentation(self) -> None:
        p = Presentation()
        p.slide_width = PInches(13.333)
        p.slide_height = PInches(7.5)

        s = self.slide_base(p, "Qual é a probabilidade de alfabetização?", 1)
        self.textbox(
            s,
            "Critério observado: proficiência ≥743 pontos\nEstimativa condicionada ao contexto do aluno",
            0.65,
            2.45,
            11.8,
            1.9,
            34,
        )
        self.wide(
            s,
            "Fatores associados · territórios · semelhanças · metas\nUma evolução da engenharia de dados da Fase 2",
            5.15,
            23,
            GOLD,
        )

        s = self.slide_base(p, "Uma base auditada e uma pipeline reproduzível", 2)
        self.cards(
            s,
            [
                ("2023", "1,50 milhão de avaliações\npara desenvolvimento"),
                ("3 folds", "Municípios inteiros\nem treino ou validação"),
                ("2024", "1,85 milhão de avaliações\npara teste separado"),
            ],
        )
        self.wide(
            s,
            "Gold por aluno: Silver e originais da Fase 2 + IBGE; nota fora dos preditores.\n"
            "Imputação e encoding dentro da pipeline, aprendidos somente no treino.\n"
            "Modelo e limiar congelados antes do teste de 2024.",
        )

        s = self.slide_base(p, "EDA: conhecer a população e o contexto", 3)
        picture = s.shapes.add_picture(
            str(ROOT / "images/16_eda_inep_2023.png"),
            PInches(0.65),
            PInches(2.15),
            width=PInches(12),
        )
        picture.left = int((p.slide_width - picture.width) / 2)

        s = self.slide_base(p, "O modelo distingue risco, com pouca seletividade", 4)
        self.cards(
            s,
            [
                ("0,516", "AP do Gradient Boosting\nno teste de 2024"),
                ("0,402", "AP da referência\nconstante"),
                ("0,622", "ROC-AUC no teste;\ngeneralização limitada"),
            ],
        )
        self.wide(
            s,
            "Comparados: baseline, logística e boosting. AP em municípios novos: 0,456.\n"
            "F2: recupera 99,30% dos casos e sinaliza 96,84% dos alunos.\n"
            "AP não é acurácia; o limiar não sustenta triagem individual autônoma.",
        )

        s = self.slide_base(p, "Inep: mais contexto, ganho preditivo pequeno", 5)
        self.cards(
            s,
            [
                ("6 → 9", "Atributos: acrescentar\nturmas, formação e jornada"),
                (">99,98%", "Cobertura dos alunos\nno desenvolvimento"),
                ("+0,000531", "Ganho médio de AP;\nmelhora em 2 dos 3 folds"),
            ],
        )
        self.wide(
            s,
            "Mesmos alunos, folds municipais e parâmetros em 2023; EDA antes dos ajustes.\n"
            "Estudo exploratório, sem novo teste independente ou significância demonstrada.\n"
            "A expansão não substitui o modelo final de seis atributos.",
        )

        s = self.slide_base(p, "Fatores associados e variáveis influentes", 6)
        self.textbox(s, "Associações com alfabetização", 0.7, 2.15, 5.7, 0.5, 24, GOLD, True)
        self.textbox(
            s,
            "Funções docentes com superior: +0,204\nAlunos por turma: −0,155\nHoras de aula: −0,025",
            0.7,
            2.85,
            5.75,
            2.1,
            22,
        )
        self.textbox(
            s,
            "Inep · Pearson em 2023\nUma observação por município × rede",
            0.7,
            5.1,
            5.75,
            0.9,
            18,
            GRAY,
        )
        self.textbox(
            s, "Influência no modelo original", 6.85, 2.15, 5.75, 0.75, 24, GOLD, True
        )
        self.textbox(
            s,
            "UF: queda de AP de 0,114826\nServiços públicos/VAB: 0,004855\nPopulação: 0,004705\nPIB per capita: 0,003733",
            6.85,
            2.85,
            5.7,
            2.1,
            22,
        )
        self.textbox(
            s,
            "Permutação em 2023 · seis atributos\nMaior queda = maior dependência preditiva",
            6.85,
            5.1,
            5.75,
            0.9,
            18,
            GRAY,
        )
        self.textbox(
            s,
            "Associação não demonstra impacto causal; correlação não é importância do modelo.",
            0.7,
            6.3,
            12,
            0.5,
            20,
            NAVY,
            True,
        )

        s = self.slide_base(p, "Municípios de maior risco e regiões semelhantes", 7)
        self.cards(
            s,
            [
                ("67,81%", "Aracaju/SE\n3.664 avaliações"),
                ("66,95%", "Nossa Senhora do Socorro/SE\n1.785 avaliações"),
                ("0,381", "Centro-Oeste e Sul:\nperfis mais próximos"),
            ],
        )
        self.wide(
            s,
            "Risco médio de não alfabetização previsto em 2024 · n ≥100 · 10 maiores em Sergipe.\n"
            "Semelhança econômica em 2023: distância padronizada, com igual peso municipal.\n"
            "Não é ranking oficial, cluster de alunos ou evidência de alfabetização igual.",
        )

        s = self.slide_base(p, "Metas: cenário de 2024 e previsão futura", 8)
        self.cards(
            s,
            [
                (
                    "1.592 / 2.819",
                    "Abaixo da meta municipal de 2024,\nentre municípios com meta disponível",
                ),
                ("2.767 / 2.924", "Abaixo de 80% no cenário\nque mantém a composição de 2024"),
            ],
        )
        self.wide(
            s,
            "Média ponderada de P(alfabetizado), municípios com ≥100 avaliações.\n"
            "Não é previsão de 2030 nem probabilidade de descumprimento de meta.\n"
            "Projetar o futuro exige novos ciclos, atributos anteriores e teste independente.",
            color=NAVY,
        )

        s = self.slide_base(p, "Demonstrar a probabilidade e interpretar a decisão", 9)
        self.cards(
            s,
            [
                ("58,66%", "P(alfabetizado): perfil\nBelo Horizonte / Municipal"),
                ("Referência 0,5", "Classificação prevista:\nalfabetizado"),
                ("Política F2", "Sinaliza atenção\ncom a mesma probabilidade"),
            ],
        )
        self.wide(
            s,
            "Demonstração histórica do modelo 1.0, com seis atributos contextuais.\n"
            "O limiar muda a decisão; não muda a probabilidade nem o critério de 743 pontos.\n"
            "O exemplo não é uma previsão individual para 2026.",
        )

        s = self.slide_base(p, "Decisões apoiadas por evidências e próximos passos", 10)
        self.cards(
            s,
            [
                ("Investigar", "Combinar risco, volume,\ncobertura e evidência local"),
                ("Ampliar", "Novos ciclos e atributos\nconhecidos antes da previsão"),
                ("Validar", "Reservar novo teste\npara projeções futuras"),
            ],
        )
        self.wide(s, "Código, métodos, resultados e limites documentados:", 5.35, 23, NAVY)
        self.textbox(
            s,
            "github.com/andreferraz-fiap-fase2/tech-challenge-fase3",
            0.65,
            5.95,
            12,
            0.55,
            21,
            GOLD,
            True,
            link=REPOSITORY,
        )
        p.core_properties.author = AUTHOR
        p.core_properties.last_modified_by = AUTHOR
        p.core_properties.title = "Alfabetização: inteligência analítica para decisões territoriais"
        p.save(self.output / f"Apresentacao-Executiva-Fase3-v{self.version}.pptx")
        script = video_script(self.version)
        (ROOT / "docs/Roteiro-Video-Fase3.md").write_text(script, encoding="utf-8")
        (self.output / f"Roteiro-Video-Fase3-v{self.version}.md").write_text(
            script, encoding="utf-8"
        )
        WordDocuments(self.output, self.version).word(
            script, f"Roteiro-Video-Fase3-v{self.version}.docx"
        )
        (self.output / "narracao.json").write_text(
            json.dumps(
                [{"slide": index, "text": text} for index, text in enumerate(NARRATIONS, 1)],
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
