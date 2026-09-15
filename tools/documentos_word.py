"""Autor: André Mohallem Ferraz. Documentos Word da entrega."""

import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

AUTHOR = "André Mohallem Ferraz"
ROOT = Path(__file__).resolve().parents[1]
NAVY, GOLD, GRAY, WHITE = "183650", "B78C31", "596573", "FFFFFF"


def markdown_paragraph_lines(markdown: str) -> list[str]:
    """Reúne quebras suaves sem quebrar ênfase, tabelas ou comandos cercados por fences."""
    lines: list[str] = []
    previous_plain, code = False, False
    for raw in markdown.splitlines():
        line = raw.strip()
        fence = line.startswith("```")
        structural = (
            not line
            or code
            or fence
            or bool(re.match(r"^(#{1,6}\s|[-*]\s|\d+\.\s|[|!>]|<!--)", line))
        )
        if previous_plain and not structural:
            lines[-1] += " " + line
        else:
            lines.append(line)
        previous_plain = not structural
        if fence:
            code = not code
    return lines


class WordDocuments:
    def __init__(self, output: Path, version: str = "1.0") -> None:
        self.output = output
        if not re.fullmatch(r"\d+\.\d+", version):
            raise ValueError(f"Versão {version!r}; esperado formato numérico N.N")
        self.version = version
        output.mkdir(parents=True, exist_ok=True)

    def inline(self, paragraph: Paragraph, text: str) -> None:
        pattern = r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
        for part in re.split(pattern, text):
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if link:
                url = link[2]
                if not url.startswith("http"):
                    url = (
                        "https://github.com/andreferraz-fiap-fase2/tech-challenge-fase3/blob/main/"
                        + url
                    )
                node = OxmlElement("w:hyperlink")
                node.set(
                    qn("r:id"),
                    paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True),
                )
                run = OxmlElement("w:r")
                properties = OxmlElement("w:rPr")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), NAVY)
                properties.append(color)
                run.append(properties)
                t = OxmlElement("w:t")
                t.text = link[1]
                run.append(t)
                node.append(run)
                paragraph._p.append(node)
            else:
                run = paragraph.add_run(
                    part[2:-2]
                    if part.startswith("**")
                    else part[1:-1]
                    if part.startswith("`")
                    else part
                )
                run.bold = part.startswith("**")
                if part.startswith("`"):
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)

    def word(self, markdown: str, filename: str) -> None:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        normal = document.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.04
        for name in ("Title", "Heading 1", "Heading 2"):
            document.styles[name].font.color.rgb = RGBColor.from_string(NAVY)
        document.styles["Title"].font.size = Pt(25)
        document.styles["Heading 1"].font.size = Pt(15)
        header = section.header.paragraphs[0]
        header.text = "FIAP  /  TECH CHALLENGE FASE 3"
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string(GOLD)
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        footer.add_run(AUTHOR + f" · v{self.version} · 14/09/2026  |  ").font.size = Pt(8)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        lines = markdown_paragraph_lines(markdown)
        i = 0
        pending_page_break = False
        code_block = False
        while i < len(lines):
            line = lines[i].strip()
            i += 1
            if not line:
                continue
            if line.startswith("```"):
                code_block = not code_block
                continue
            if code_block:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                if pending_page_break:
                    paragraph.paragraph_format.page_break_before = True
                    pending_page_break = False
                continue
            if line == "<!-- pagebreak -->":
                pending_page_break = True
                continue
            image = re.fullmatch(r"!\[.*?\]\((.*?)\)", line)
            if image:
                width = 6.0 if image[1].endswith("02_regioes_2023.png") else 6.65
                document.add_picture(str(ROOT / image[1]), width=Inches(width))
                document.paragraphs[-1].alignment = 1
                continue
            if line.startswith("|"):
                rows = [line]
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append(lines[i].strip())
                    i += 1
                parsed = [[c.strip() for c in row.strip("|").split("|")] for row in rows]
                parsed = [row for row in parsed if not all(re.fullmatch(r":?-+:?", c) for c in row)]
                table = document.add_table(rows=0, cols=len(parsed[0]))
                table.style = "Table Grid"
                for index, row in enumerate(parsed):
                    cells = table.add_row().cells
                    for cell, text in zip(cells, row):
                        self.inline(cell.paragraphs[0], text)
                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(9)
                        if index == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                                run.font.color.rgb = RGBColor.from_string(WHITE)
                            shade = OxmlElement("w:shd")
                            shade.set(qn("w:fill"), NAVY)
                            cell._tc.get_or_add_tcPr().append(shade)
                    properties = table.rows[index]._tr.get_or_add_trPr()
                    properties.append(OxmlElement("w:cantSplit"))
                    if index == 0:
                        properties.append(OxmlElement("w:tblHeader"))
                continue
            if line.startswith("### "):
                paragraph = document.add_paragraph(style="Heading 2")
                text = line[4:]
            elif line.startswith("# "):
                paragraph = document.add_paragraph(style="Title")
                text = line[2:]
            elif line.startswith("## "):
                paragraph = document.add_paragraph(style="Heading 1")
                text = line[3:]
            elif line.startswith("- "):
                paragraph = document.add_paragraph(style="List Bullet")
                text = line[2:]
            else:
                paragraph = document.add_paragraph()
                text = line
            if pending_page_break:
                paragraph.paragraph_format.page_break_before = True
                pending_page_break = False
            self.inline(paragraph, text)
        document.core_properties.author = AUTHOR
        document.core_properties.last_modified_by = AUTHOR
        document.core_properties.title = markdown.splitlines()[0].lstrip("# ")
        document.core_properties.created = datetime(2026, 9, 14, tzinfo=timezone.utc)
        document.core_properties.modified = datetime(2026, 9, 14, tzinfo=timezone.utc)
        document.save(self.output / filename)

    def report(self) -> None:
        readme = (ROOT / "README.md").read_text()
        pieces = re.split(r"(?=^## )", readme, flags=re.M)
        sections = {piece.splitlines()[0]: piece for piece in pieces[1:]}
        title = (
            "# Relatório técnico — alfabetização no Brasil\n\n**Autor: "
            + AUTHOR
            + f"**\n\n**FIAP · Tech Challenge Fase 3 · Trabalho individual · Revisão documental {self.version} · 14/09/2026**\n\n"
            + "**Repositório do projeto:** [github.com/andreferraz-fiap-fase2/tech-challenge-fase3](https://github.com/andreferraz-fiap-fase2/tech-challenge-fase3).\n\n"
        )
        abstract = "**Síntese executiva.** O Gradient Boosting supera o baseline no teste temporal, com AP 0,5162 e ROC-AUC 0,6224. Seu limiar acadêmico de F2 sinaliza 96,84% dos alunos. A entrega evidencia potencial para leitura territorial e limites importantes de generalização e seletividade, sem recomendar decisões individuais autônomas.\n\n"
        groups = [
            ["## 1. Contexto do problema", "## 2. Objetivo analítico"],
            ["## 3. Base utilizada"],
            ["## 4. Análise exploratória e entendimento do problema"],
            ["## 5. Etapas de modelagem"],
            ["## 6. Escolha do algoritmo"],
            ["## 7. Métricas de avaliação"],
            ["## 8. Interpretação dos resultados"],
            ["## 9. Insights encontrados"],
            ["## 10. Limitações", "## 11. Aplicação prática para políticas públicas"],
            ["## Demonstração da previsão"],
            ["## 12. Possíveis evoluções futuras", "## Como reproduzir"],
            ["## Referências"],
        ]
        pages = []
        extras = {
            4: "\n![Escolha do limiar](images/09_limiar_f2_2023.png)\n",
            7: "\n![Erros por região](images/12_regioes_teste_2024.png)\n",
            11: "\n## Apêndice — perfis regionais\n\n![Perfis regionais](images/14_perfis_regionais_2023.png)\n",
        }
        for index, group in enumerate(groups):
            page = "\n".join(sections[name] for name in group)
            if index == 2:
                page = page.replace("### 4.2.", "<!-- pagebreak -->\n\n### 4.2.")
                page = page.replace("### 4.3.", "<!-- pagebreak -->\n\n### 4.3.")
                page = page.replace("### 4.4.", "<!-- pagebreak -->\n\n### 4.4.")
            if index == 3:
                page = page.replace("### 5.2.", "<!-- pagebreak -->\n\n### 5.2.")
            if index == 0:
                page = page.replace("## 2.", "<!-- pagebreak -->\n\n## 2.")
            if index == 10:
                page = re.sub(r"```text.*?```", "", page, flags=re.S)
            pages.append(page + extras.get(index, ""))
        pages.append(
            "## Apêndice — distribuições do contexto\n\nUma observação por município no desenvolvimento de 2023. A escala log10 desta figura é apenas visual; o modelo logístico utiliza log1p.\n\n![Distribuições do contexto](images/03_contexto_municipal_2023.png)\n"
        )
        pages.append(
            "## Apêndice — linhagem da Gold\n\nA Gold por aluno adapta a granularidade das fontes da Fase 2. A Gold municipal original participa da análise posterior de metas. O ramo educacional utiliza somente desenvolvimento e não substitui o modelo temporalmente avaliado.\n\n![Linhagem da Gold](images/15_linhagem_gold.png)\n\n[Detalhes da pergunta e das origens](docs/Pergunta-e-Linhagem.md).\n"
        )
        markdown = title + abstract + "\n<!-- pagebreak -->\n".join(pages)
        (ROOT / "docs/Relatorio-Tecnico-Fase3.md").write_text(
            re.sub(r"\]\((?!https?://)([^)]+)\)", lambda match: "](../" + match[1] + ")", markdown)
        )
        filename = f"Relatorio-Tecnico-Fase3-v{self.version}"
        self.word(markdown, filename + ".docx")
        (self.output / (filename + ".md")).write_text(markdown)
